import os
import sys
from fpdf import FPDF
from docx import Document

def convert_txt_to_pdf(txt_file, output_pdf):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    with open(txt_file, "r", encoding="utf-8") as file:
        for line in file:
            pdf.cell(200, 10, txt=line.strip(), ln=True)
    
    pdf.output(output_pdf)
    print(f"PDF gerado: {output_pdf}")

def convert_txt_to_doc(txt_file, output_doc):
    doc = Document()
    
    with open(txt_file, "r", encoding="utf-8") as file:
        for line in file:
            doc.add_paragraph(line.strip())
    
    doc.save(output_doc)
    print(f"DOC gerado: {output_doc}")

def main():
    if len(sys.argv) < 4:
        print("Uso: python script.py <arquivo.txt> <formato: pdf/doc> <arquivo_saida>")
        sys.exit(1)
    
    input_txt = sys.argv[1]
    output_format = sys.argv[2].lower()
    output_file = sys.argv[3]
    
    if not os.path.exists(input_txt):
        print("Erro: Arquivo TXT não encontrado!")
        sys.exit(1)
    
    if output_format == "pdf":
        convert_txt_to_pdf(input_txt, output_file)
    elif output_format == "doc":
        convert_txt_to_doc(input_txt, output_file)
    else:
        print("Erro: Formato inválido! Use 'pdf' ou 'doc'")
        sys.exit(1)

if __name__ == "__main__":
    main()
